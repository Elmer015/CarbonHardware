"""Build the CPU/GPU STEC research workbook (.xlsx) from repository data only.

Run from the repository root:

    python ml-engine/carbon/stec_batch.py
    python ml-engine/experiments/build_stec_workbook.py

Inputs: data/processed/stec-per-chip.csv (batch output), data/stec/*.csv (evidence),
data/raw/grid/ (OWID carbon intensity), CarbonSet splits (model validation).
Output: data/processed/Analisis_STEC_CPU_GPU.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "ml-engine" / "carbon"))
sys.path.insert(0, str(ROOT / "ml-engine" / "experiments"))

from cci import compute_cci
from operational_carbon import calculate_from_tdp
from stec import NODE_PARAMETERS, calculate_stec
from stec_evidence import EVIDENCE_DIR, GRID_PATH, GRID_SOURCE, load_evidence, resolve_processor_stec
import type_validation

BATCH_CSV = ROOT / "data" / "processed" / "stec-per-chip.csv"
BATCH_SUMMARY = ROOT / "data" / "processed" / "stec-batch-summary.json"
CLEANED = ROOT / "data" / "interim" / "carbonset-cleaned.csv"
DEFAULT_OUTPUT = ROOT / "data" / "processed" / "Analisis_STEC_CPU_GPU.xlsx"
PAPER = "Zhang, Yang, Wang - arXiv:2312.06364v1 (ACM e-Energy 2024)"
PAPER_TC = "IEEE Trans. Comput. 74(12):4037-4049, 2025, doi 10.1109/TC.2025.3605743"
PAPER_TABLE5 = {"Taiwan": 561, "China": 531, "South Korea": 436, "United States": 367, "Italy": 372, "Ireland": 346}
GRID_COUNTRIES = ["Taiwan", "South Korea", "United States", "Ireland", "Israel", "China", "Germany", "Italy", "Indonesia", "World"]


def reproduction_checks() -> pd.DataFrame:
	rows = []
	for node in ("28", "20", "14", "10", "7", "7-EUV", "5", "3"):
		for country, value in PAPER_TABLE5.items():
			r = calculate_stec(die_area_mm2=100, node_nm=node, grid_intensity_g_co2e_per_kwh=value, manufacturing_region=country, manufacturing_year=2022, source=f"{PAPER} Table 5")
			p = NODE_PARAMETERS[node]
			rows.append({
				"Pemeriksaan": "Eq.(7) per cm2 dengan CI Tabel 5 (2022)", "Node": node, "Negara": country, "CI (g/kWh)": value,
				"GPS (g/cm2)": p.gps_g_per_cm2, "MPS (g/cm2)": p.mps_g_per_cm2, "Listrik (g/cm2)": value * p.eps_kwh_per_cm2,
				"Total (g/cm2)": r.intensity_g_per_cm2, "Porsi listrik": value * p.eps_kwh_per_cm2 / r.intensity_g_per_cm2,
			})
	return pd.DataFrame(rows)


def implied_ci_checks() -> pd.DataFrame:
	rows = []
	for label, total in (("STEC-ZY EU", 1266.88), ("STEC-ZY ASEAN", 1848.10), ("STEC-GY global", 1557.49)):
		for node in ("7", "7-EUV"):
			p = NODE_PARAMETERS[node]
			rows.append({"Nilai paper (Tabel 4, CPU 7nm)": label, "Total paper (g/cm2)": total, "Parameter node": node,
				"CI tersirat (g/kWh)": (total - p.gps_g_per_cm2 - p.mps_g_per_cm2) / p.eps_kwh_per_cm2,
				"Catatan": "CI tersirat masuk akal untuk zona tsb jika node 7-EUV; dengan node 7 (DUV) ASEAN tersirat > 700 g/kWh"})
	p = NODE_PARAMETERS["7-EUV"]
	rows.append({"Nilai paper (Tabel 4, CPU 7nm)": "Porsi Sec. 2.3: GPS 11%, MPS 28%, EPS 61%", "Total paper (g/cm2)": p.mps_g_per_cm2 / 0.28,
		"Parameter node": "7-EUV", "CI tersirat (g/kWh)": (p.mps_g_per_cm2 / 0.28 * 0.61) / p.eps_kwh_per_cm2,
		"Catatan": "porsi paper konsisten dengan parameter 7-EUV dan CI ~ 500 g/kWh"})
	return pd.DataFrame(rows)


def example_system(cleaned: pd.DataFrame) -> pd.DataFrame:
	evidence = load_evidence()
	use_country, use_year = "Indonesia", 2024
	ci_use = evidence.grid[(use_country, use_year)]
	scenario = {"utilization": 0.5, "hours_per_day": 8.0, "days_per_year": 365.0, "lifetime_years": 4.0}
	service_hours = scenario["hours_per_day"] * scenario["days_per_year"]

	def device(name):
		row = cleaned[cleaned["Product"].eq(name)].iloc[0]
		stec = resolve_processor_stec({"name": name, "type": row["Type"], "vendor": row["Vendor"], "processSizeNm": row["Process Size (nm)"],
			"dieAreaMm2": row["DieSizeValue"], "releaseYear": int(str(row["Release Date"])[:4])}, evidence=evidence)
		op = calculate_from_tdp(row["TDP (W)"], scenario["utilization"], scenario["hours_per_day"], scenario["days_per_year"],
			ci_use / 1000, grid_intensity_source=GRID_SOURCE, analysis_period=str(use_year), assumptions={"setting": "workstation, PUE tidak berlaku"})
		return {"name": name, "stec": stec, "op": op.carbon_kg_co2e, "carbonset": row["Emb_Carbon_Mean"]}

	devices = {"CPU": device("AMD Ryzen 9 5950X"), "GPU": device("AMD Radeon RX 6800 XT")}
	rows = []
	for config, members in (("CPU saja", ["CPU"]), ("GPU saja", ["GPU"]), ("CPU + GPU", ["CPU", "GPU"])):
		emb = {k: devices[k]["stec"]["stecKgCo2e"] if k in members else 0.0 for k in ("CPU", "GPU")}
		op = {k: devices[k]["op"] if k in members else 0.0 for k in ("CPU", "GPU")}
		annual_emb = (emb["CPU"] + emb["GPU"]) / scenario["lifetime_years"]
		annual_total = annual_emb + op["CPU"] + op["GPU"]
		cci = compute_cci(annual_total, service_hours, workload_unit="1 jam pemakaian workstation pada utilisasi 50%",
			lifecycle_boundary="STEC die (manufaktur) + listrik pemakaian", analysis_period=f"1 tahun layanan ({use_year})")
		rows.append({
			"Konfigurasi": config, "Perangkat": " + ".join(devices[m]["name"] for m in members),
			"Embodied CPU - STEC (kg, total)": emb["CPU"], "Embodied GPU - STEC (kg, total)": emb["GPU"],
			"Embodied teramortisasi (kg/tahun)": annual_emb, "Operasional CPU (kg/tahun)": op["CPU"], "Operasional GPU (kg/tahun)": op["GPU"],
			"Total karbon (kg/tahun)": annual_total, "Porsi embodied": annual_emb / annual_total,
			"CCI (kg CO2e per jam pemakaian)": cci.cci_kg_co2e_per_workload,
			"Referensi CarbonSet Emb_Carbon_Mean (kg, TIDAK dijumlahkan)": sum(devices[m]["carbonset"] for m in members),
			"Level bukti lokasi": " / ".join(devices[m]["stec"]["location"]["level"] for m in members),
		})
	rows.append({"Konfigurasi": "ASUMSI", "Perangkat": (
		f"STEC dari resolver evidence (5950X = chiplet CCD Taiwan + cIOD AS; RX 6800 XT = asumsi vendor/node Taiwan); "
		f"tahun manufaktur = tahun rilis (proxy); pemakaian: {use_country} {use_year}, CI {ci_use:.1f} g/kWh (OWID); "
		f"utilisasi {scenario['utilization']}, {scenario['hours_per_day']:.0f} jam/hari, {scenario['days_per_year']:.0f} hari/tahun, "
		f"lifetime {scenario['lifetime_years']:.0f} tahun; daya = TDP x utilisasi (proxy).")})
	return pd.DataFrame(rows)


def methodology_rows() -> pd.DataFrame:
	R = []
	def add(topic, decision, source):
		R.append({"Topik": topic, "Keputusan / definisi": decision, "Sumber / status": source})
	add("Persamaan STEC (CPU/GPU die)", "EC_C(s,t) = GPS + MPS + CI(s,t) x EPS [g CO2e/cm2], x luas die (cm2) / 1000 -> kg CO2e. Yield opsional: dibagi yield_factor sesuai Eq.(1)", f"{PAPER} Eq.(1),(7); {PAPER_TC}")
	add("Carbon intensity CI(s,t)", "Tabel negara-tahun (STEC-CY) dari OWID/Ember; tidak dihitung ulang dari bauran energi Eq.(2)", f"Eq.(2) paper; {GRID_PATH.relative_to(ROOT)}")
	add("GPS / MPS / EPS", "GPS & MPS dalam g CO2e/cm2, EPS dalam kWh/cm2 (header Tabel 7 tertulis g/cm2, tetapi Tabel 9 mendefinisikan kWh/cm2)", "Tabel 7 & 9 paper")
	add("Proxy node", "22nm -> parameter 20nm; 12nm (GF 12LP, IO-die) -> parameter 14nm. Keduanya berlabel explicit_proxy_node", "Keputusan metodologi (stec.py)")
	add("Atribusi spasial - urutan bukti", "1) input pengguna, 2) fab-provenance.csv (sumber per produk), 3) chiplets.csv (per die), 4) fab-rules.csv (asumsi vendor/node; >1 kandidat -> rentang), 5) unresolved", "stec_evidence.py; data/stec")
	add("Atribusi temporal", "Tahun manufaktur: input pengguna > manufacturing-years.csv > tahun rilis CarbonSet (proxy berlabel). Rilis =/= manufaktur =/= deploy =/= retirement", "stec_evidence.py")
	add("Dual attribution", "Location-based: CI grid lokasi & tahun fab. Market-based: hanya dari market-factors.csv (faktor kontraktual + sumber); tanpa itu status 'not_available'. Tidak pernah dijumlahkan", f"Abstrak {PAPER_TC}; stec.calculate_market_based_stec")
	add("Chiplet / multi-die", "Setiap die dihitung dengan luas, node, region, dan CI sendiri lalu dijumlahkan (calculate_multi_die_stec). CarbonSet hanya mencatat CCD untuk chiplet AMD; IO-die tidak tercatat", "chiplets.csv (ISSCC 2020 dsb.)")
	add("Yield", "Hanya diterapkan jika ada di yields.csv (node, foundry, tahun, sumber). Default: tidak dikoreksi yield", "stec_evidence.py")
	add("Boundary", "Cradle-to-gate die saja. Tidak termasuk packaging, substrate, PCB, HBM/memori, cooler, transport, use, end-of-life", "Konsisten dengan model CPU paper")
	add("Functional unit", "Embodied: kg CO2e per chip (die saja). Operasional: kg CO2e per tahun. CCI: kg CO2e per unit workload yang didefinisikan", "cci.py, stec.py")
	add("Anti double counting", "Jangan jumlahkan STEC dengan CarbonSet Emb_Carbon_*; jangan jumlahkan CarbonSet Ope_Carbon_* dengan operational_carbon.py; location- dan market-based dilaporkan terpisah", "Keputusan metodologi")
	add("Alokasi multi-workload", "Embodied per workload = STEC x (jam workload / total jam layanan selama L)", "Keputusan metodologi")
	add("Perbedaan basis CI", "Tabel 5 paper (2022): Taiwan 561 g/kWh; seri OWID saat ini: Taiwan 2022 = 639 g/kWh (basis berbeda)", "Ditemukan saat implementasi - nyatakan basis di paper")
	return pd.DataFrame(R)


def input_contract() -> pd.DataFrame:
	fields = [
		("inputMode", "existing | custom", "-", "wajib", "API /api/estimates/stec/processor", "400"),
		("processorName", "teks", "-", "wajib", "lookup CarbonSet / evidence", "400"),
		("processSizeNm", "angka", "nm", "wajib (custom)", "parameter Tabel 7", "STECIncompleteError -> 400"),
		("dieSizeMm2", "angka (total semua die)", "mm2", "wajib (custom)", "STEC", "400"),
		("tdpW", "angka", "W", "wajib untuk operasional", "proxy daya", "400"),
		("manufacturingRegion", "negara (nama OWID)", "-", "opsional (existing), wajib (custom)", "CI(s,t)", "existing: pakai evidence; custom: 400"),
		("manufacturingYear", "tahun", "tahun", "opsional (existing), wajib (custom)", "CI(s,t)", "existing: tahun rilis (proxy berlabel)"),
		("fab-provenance.csv", "vendor,type,product,process_node_nm,manufacturing_region,foundry,confidence,source", "-", "opsional", "titik per produk", "jatuh ke fab-rules (asumsi)"),
		("chiplets.csv", "product,die_role,die_count,die_area_mm2,node_nm,manufacturing_region,foundry,confidence,source", "mm2", "opsional", "STEC multi-die", "luas die CarbonSet dipakai"),
		("manufacturing-years.csv", "product,release_year,manufacturing_year,source", "tahun", "opsional", "tahun manufaktur", "tahun rilis (proxy)"),
		("yields.csv", "process_node_nm,foundry,year,yield_factor,source", "0-1", "opsional", "koreksi yield", "tidak dikoreksi"),
		("market-factors.csv", "foundry,year,market_based_factor_g_co2e_per_kwh,contract_type,source", "g CO2e/kWh", "opsional", "STEC market-based", "status not_available"),
		("utilization, hours/day, days/year, lifetime", "angka", "-, jam, hari, tahun", "wajib untuk operasional & CCI", "operasional, amortisasi", "400"),
		("grid factor pemakaian + sumber", "angka + sitasi", "kg CO2e/kWh", "wajib untuk operasional carbon", "operasional", "carbonStatus pending_grid_intensity"),
		("workload", "unit + jumlah", "unit kerja", "wajib untuk CCI", "CCI", "400"),
	]
	return pd.DataFrame(fields, columns=["Field / file", "Tipe / kolom", "Satuan", "Wajib?", "Dipakai untuk", "Jika tidak tersedia"])


def status_rows(validation: dict, summary: dict) -> pd.DataFrame:
	m = validation["metrics"]
	def get(typ):
		row = m[(m.model == "A_saved_artifact_no_type") & (m.split == "test") & (m.type == typ)].iloc[0]
		return f"R2 {row.r2:.3f}, MAE {row.mae:.3f} kg"
	d = validation["die_cv"]
	def die(typ):
		row = d[(d.model == "DieGroupKFold5_no_type") & (d.type == typ)].iloc[0]
		return f"R2 {row.r2:.3f}, MAE {row.mae:.2f} kg, MAPE {row.mape_pct:.1f}%"
	leak = validation["leakage"].set_index("check")["value"]
	sc, lc = summary["status_counts"], summary["location_level_counts"]
	batch = f"{summary['processors']} prosesor: {sc.get('point', 0)} titik, {sc.get('range', 0)} rentang, {sc.get('incomplete', 0)} incomplete"
	return pd.DataFrame([
		("Batch STEC_per_Chip", "SELESAI", batch, "ml-engine/carbon/stec_batch.py -> data/processed/stec-per-chip.csv", f"DATA: data/interim/carbonset-cleaned.csv + data/stec + {GRID_PATH.relative_to(ROOT)}"),
		("Workbook di repository", "SELESAI", "Workbook dibangun ulang hanya dari data di repo", "ml-engine/experiments/build_stec_workbook.py -> data/processed/Analisis_STEC_CPU_GPU.xlsx", "DATA: output batch + evidence + split CarbonSet"),
		("Fab provenance per produk", "SEBAGIAN", f"{lc.get('product_sourced', 0)} produk bersumber + {summary['chiplet_products']} produk chiplet; {lc.get('vendor_node_assumption', 0)} masih asumsi vendor/node", "data/stec/fab-provenance.csv, chiplets.csv, fab-rules.csv", "DATA: whitepaper NVIDIA A100/GA102/V100; ISSCC 2020 (AMD chiplet). ASUMSI: fab-rules.csv"),
		("Tahun manufaktur aktual", "MEKANISME SIAP, DATA BELUM", f"{summary['manufacturing_year_status_counts']}", "data/stec/manufacturing-years.csv (header saja)", "Belum ada data bersumber - tahun rilis dipakai sebagai proxy berlabel"),
		("Yield per node/fab", "MEKANISME SIAP, DATA BELUM", f"{summary['yield_corrected_products']} produk dikoreksi yield", "data/stec/yields.csv (header saja)", "Belum ada data bersumber"),
		("Chiplet multi-node", "SELESAI (4 produk)", f"{summary['chiplet_products']} produk dimodelkan per die (CCD + IO-die)", "stec.calculate_multi_die_stec + chiplets.csv", "DATA: Naffziger et al. ISSCC 2020; Zen 3 CCD (Hot Chips/ISSCC 2021) - perlu verifikasi"),
		("Market-based attribution", "MEKANISME SIAP, DATA BELUM", f"{summary['market_based_computed_products']} produk dihitung market-based", "stec.calculate_market_based_stec + market-factors.csv (header saja)", "Belum ada faktor bersumber - butuh laporan ESG foundry (PPA/REC)"),
		("Backend/frontend integration", "SELESAI", "Form -> POST /api/estimates/stec/processor -> Python resolver -> kartu hasil STEC di UI", "backend/src (controller, route, python.service), frontend/src (HardwareForm, StecResult)", "Diuji: node --test (3 lulus), vite build, panggilan API end-to-end"),
		("Validasi model CPU vs GPU", "SELESAI", f"Test (split produk): CPU {get('CPU')}; GPU {get('GPU')}", "ml-engine/experiments/type_validation.py", "DATA (dihitung): split CarbonSet + artifact Phase 4"),
		("Kebocoran data", "TEMUAN KRITIS", f"{leak['test rows whose die signature also appears in train']} dari {leak['test rows']} baris test memakai die yang juga ada di train", "type_validation.py", "DATA (dihitung)"),
		("Validasi jujur (split per die)", "HASIL", f"CPU {die('CPU')}; GPU {die('GPU')}", "type_validation.py", "DATA (dihitung)"),
		("Kualitas data CarbonSet", "TEMUAN", "Chiplet AMD hanya mencatat CCD (EPYC 7742 tercatat 1x74 mm2); RTX 3090 Ti bertanggal rilis 2020-01-01", "STEC_per_Chip, chiplets.csv", "DATA: CarbonSet.csv dibandingkan dengan sumber chiplet/whitepaper"),
	], columns=["Item", "Status", "Keterangan", "Lokasi di repo", "Data / sumber"])


def style_workbook(writer: pd.ExcelWriter) -> None:
	header_fill = PatternFill("solid", fgColor="0D1B3E")
	for sheet in writer.book.worksheets:
		sheet.freeze_panes = "A2"
		for cell in sheet[1]:
			cell.font = Font(bold=True, color="FFFFFF")
			cell.fill = header_fill
			cell.alignment = Alignment(wrap_text=True, vertical="top")
		for idx, column in enumerate(sheet.columns, start=1):
			values = [str(c.value) if c.value is not None else "" for c in column]
			sheet.column_dimensions[get_column_letter(idx)].width = max(min(max(len(v) for v in values) + 2, 70), 10)
			for cell in column[1:]:
				if isinstance(cell.value, float):
					cell.number_format = "0.000"
				if isinstance(cell.value, str) and len(cell.value) > 60:
					cell.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
	parser = argparse.ArgumentParser(description=__doc__)
	parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
	args = parser.parse_args()
	if not BATCH_CSV.exists():
		raise SystemExit("Run ml-engine/carbon/stec_batch.py first")

	evidence = load_evidence()
	cleaned = pd.read_csv(CLEANED)
	chips = pd.read_csv(BATCH_CSV)
	summary = json.loads(BATCH_SUMMARY.read_text(encoding="utf-8"))
	validation = type_validation.run()

	grid = pd.DataFrame([{"year": y, **{c: evidence.grid.get((c, y)) for c in GRID_COUNTRIES}} for y in range(2011, 2025)])
	grid.insert(0, "Sumber", GRID_SOURCE)
	table5 = pd.DataFrame([{"Negara": k, "CI 2022 Tabel 5 paper (g/kWh)": v, "CI 2022 OWID (g/kWh)": evidence.grid.get((k, 2022))} for k, v in PAPER_TABLE5.items()])
	nodes = pd.DataFrame([{"Node": k, "Parameter node": p.node_nm, "EPS (kWh/cm2)": p.eps_kwh_per_cm2, "GPS (g/cm2)": p.gps_g_per_cm2, "MPS (g/cm2)": p.mps_g_per_cm2, "Sumber": p.source} for k, p in NODE_PARAMETERS.items()])
	ref = chips["stec_kg_co2e"].fillna(chips["global_baseline_stec_kg_co2e"])
	chips["rasio_stec_per_carbonset"] = ref / chips["carbonset_emb_mean_kg_co2e"]
	ringkasan = chips.groupby(["type", "status", "location_level"]).agg(
		jumlah=("product", "size"), median_stec_kg=("stec_kg_co2e", "median"), median_global_baseline_kg=("global_baseline_stec_kg_co2e", "median"),
		median_carbonset_kg=("carbonset_emb_mean_kg_co2e", "median"), median_rasio=("rasio_stec_per_carbonset", "median")).reset_index()
	corr = chips.groupby("type")[["global_baseline_stec_kg_co2e", "carbonset_emb_mean_kg_co2e"]].corr(method="spearman").xs("global_baseline_stec_kg_co2e", level=1)["carbonset_emb_mean_kg_co2e"]
	ringkasan["Spearman STEC-GY vs CarbonSet (per type)"] = ringkasan["type"].map(corr)

	args.output.parent.mkdir(parents=True, exist_ok=True)
	with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
		status_rows(validation, summary).to_excel(writer, sheet_name="Ringkasan_Status", index=False)
		methodology_rows().to_excel(writer, sheet_name="Metodologi_STEC", index=False)
		input_contract().to_excel(writer, sheet_name="Kontrak_Input", index=False)
		nodes.to_excel(writer, sheet_name="Parameter_Node", index=False)
		grid.to_excel(writer, sheet_name="Grid_CI_OWID", index=False)
		table5.to_excel(writer, sheet_name="Grid_CI_vs_Paper", index=False)
		for name in ("fab-provenance", "chiplets", "fab-rules", "manufacturing-years", "yields", "market-factors"):
			pd.read_csv(EVIDENCE_DIR / f"{name}.csv").to_excel(writer, sheet_name=f"Evidence_{name}"[:31], index=False)
		chips.to_excel(writer, sheet_name="STEC_per_Chip", index=False)
		ringkasan.to_excel(writer, sheet_name="STEC_Ringkasan", index=False)
		reproduction_checks().to_excel(writer, sheet_name="Reproduksi_Eq7", index=False)
		implied_ci_checks().to_excel(writer, sheet_name="Cek_Tabel4_Paper", index=False)
		example_system(cleaned).to_excel(writer, sheet_name="Contoh_Hasil_CPU_GPU", index=False)
		validation["metrics"].to_excel(writer, sheet_name="Validasi_Model_Type", index=False)
		validation["leakage"].to_excel(writer, sheet_name="Cek_Kebocoran", index=False)
		validation["die_cv"].to_excel(writer, sheet_name="Validasi_Split_Die", index=False)
		validation["lovo"].to_excel(writer, sheet_name="Validasi_LOVO", index=False)
		validation["largest_errors"].to_excel(writer, sheet_name="Error_Terbesar", index=False)
		validation["out_of_range"].to_excel(writer, sheet_name="Di_Luar_Rentang", index=False)
		pd.concat([validation["counts"], validation["ambiguous"].assign(Type="produk ambigu (dikeluarkan)")], ignore_index=True).to_excel(writer, sheet_name="Jumlah_Data", index=False)
		style_workbook(writer)
	print(f"wrote {args.output} | chips={len(chips)}")


if __name__ == "__main__":
	main()
